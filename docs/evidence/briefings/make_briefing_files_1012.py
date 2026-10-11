#!/usr/bin/env python3
"""Build the briefing PDF/DOCX with URLs on their OWN lines.

Inline long URLs inside justified paragraphs cannot be preserved by any
reportlab setting (probed: splitLongWords, nobr, alignment, wordWrap,
non-breaking hyphen — only splitLongWords=0 kept the hyphen and even then
the URL split across lines). URLs on their own line at 8pt inside <nobr>
were verified byte-perfect, so this converter emits them that way.
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate

DATE = "2026-10-12"
SRC = os.path.join(os.path.expanduser("~"), ".hermes", "cron", "output",
                   "sustainability-briefing-%s.src.md" % DATE)
BASE = SRC.replace(".src.md", "")
text = open(SRC, encoding="utf-8").read()

GREEN = RGBColor(0x1B, 0x5E, 0x20)
GREY = RGBColor(0x55, 0x55, 0x55)
URL_RE = re.compile(r"^https?://")


def esc(s):
    s = s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<i>\1</i>", s)
    return s


def clean(s):
    c = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    c = re.sub(r"\*(.+?)\*", r"\1", c)
    c = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1 (\2)", c)
    return c


# ---------- DOCX ----------
doc = Document()
n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(6)
n.paragraph_format.line_spacing = 1.15


def rich(p, s):
    for part in re.split(r'(\*\*.+?\*\*)', s):
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            p.add_run(part)


for line in text.split("\n"):
    s = line.strip()
    if not s:
        continue
    if s.startswith("# "):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(clean(s[2:])); r.bold = True; r.font.size = Pt(18)
        r.font.color.rgb = GREEN
        p.paragraph_format.space_after = Pt(4)
    elif s.startswith("## "):
        p = doc.add_paragraph()
        r = p.add_run(clean(s[3:])); r.bold = True; r.font.size = Pt(14)
        r.font.color.rgb = GREEN
        p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
    elif s == "---":
        continue
    elif re.match(r"^\*\*[A-Z][a-z]+day,", s):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(clean(s)); r.font.size = Pt(12); r.font.color.rgb = GREY
    elif URL_RE.match(s):
        p = doc.add_paragraph()
        r = p.add_run(s); r.font.size = Pt(8); r.font.color.rgb = GREY
        p.paragraph_format.space_after = Pt(1)
    elif s.startswith("- **"):
        p = doc.add_paragraph(style='List Bullet')
        rich(p, s[2:])
        p.paragraph_format.space_after = Pt(6)
    elif re.match(r"^\d+\.\s", s):
        p = doc.add_paragraph(style='List Number')
        rich(p, re.sub(r"^\d+\.\s", "", s))
        p.paragraph_format.space_after = Pt(6)
    elif s.startswith("*Sources") or s.lower().startswith("sources"):
        p = doc.add_paragraph()
        r = p.add_run(s); r.bold = True; r.font.size = Pt(9)
    elif s.startswith("*") and s.endswith("*") and not s.startswith("**"):
        p = doc.add_paragraph()
        r = p.add_run(s.strip("*")); r.italic = True; r.font.size = Pt(10)
        r.font.color.rgb = GREY
    else:
        p = doc.add_paragraph()
        rich(p, s)

doc.save(BASE + ".docx")
print("DOCX saved:", BASE + ".docx")

# ---------- PDF ----------
styles = getSampleStyleSheet()
body = ParagraphStyle('body', parent=styles['Normal'], fontName='Helvetica',
                      fontSize=10, leading=14.5, spaceAfter=7, alignment=4)
h1 = ParagraphStyle('h1', parent=styles['Normal'], fontName='Helvetica-Bold',
                    fontSize=18, leading=22, textColor=colors.HexColor('#1B5E20'),
                    alignment=1, spaceAfter=4)
datest = ParagraphStyle('date', parent=styles['Normal'], fontName='Helvetica',
                        fontSize=11, leading=14, textColor=colors.HexColor('#555555'),
                        alignment=1, spaceAfter=12)
h2 = ParagraphStyle('h2', parent=styles['Normal'], fontName='Helvetica-Bold',
                    fontSize=13, leading=16, textColor=colors.HexColor('#1B5E20'),
                    spaceBefore=14, spaceAfter=6)
bullet = ParagraphStyle('bullet', parent=body, leftIndent=14, bulletIndent=4,
                        spaceAfter=7, splitLongWords=0)
numlst = ParagraphStyle('numlst', parent=body, leftIndent=16, bulletIndent=2,
                        spaceAfter=7, splitLongWords=0)
# 8pt is the largest size at which the longest source URL fits one line
# inside the 17cm text column (measured: 437pt at 8pt vs 491.7pt at 9pt).
urlst = ParagraphStyle('url', parent=body, fontSize=8, leading=11,
                       textColor=colors.HexColor('#555555'), spaceAfter=1)
srcst = ParagraphStyle('src', parent=body, fontSize=9, leading=12,
                       textColor=colors.HexColor('#555555'), spaceAfter=3)
ital = ParagraphStyle('ital', parent=body, fontSize=9, leading=12,
                      textColor=colors.HexColor('#555555'))

story = []
num = 0
for line in text.split("\n"):
    s = line.strip()
    if not s:
        continue
    if s.startswith("# "):
        story.append(Paragraph(esc(s[2:]), h1))
    elif re.match(r"^\*\*[A-Z][a-z]+day,", s):
        story.append(Paragraph(esc(s), datest))
    elif s.startswith("## "):
        story.append(Paragraph(esc(s[3:]), h2))
    elif s == "---":
        story.append(HRFlowable(width="100%", thickness=0.6,
                                color=colors.HexColor('#CCCCCC'),
                                spaceBefore=6, spaceAfter=6))
    elif URL_RE.match(s):
        # <nobr> + a size that fits the column => byte-perfect URL
        story.append(Paragraph("<nobr>%s</nobr>" % esc(s), urlst))
    elif s.startswith("- **"):
        story.append(Paragraph(esc(s[2:]), bullet, bulletText='\u2022'))
    elif re.match(r"^\d+\.\s", s):
        num += 1
        story.append(Paragraph(esc(re.sub(r"^\d+\.\s", "", s)), numlst,
                               bulletText='%d.' % num))
    elif s.startswith("*Sources") or s.lower().startswith("sources"):
        story.append(Paragraph(esc(s), srcst))
    elif s.startswith("*") and s.endswith("*") and not s.startswith("**"):
        story.append(Paragraph(esc(s), ital))
    else:
        num = 0
        story.append(Paragraph(esc(s), body))

SimpleDocTemplate(BASE + ".pdf", pagesize=A4, leftMargin=2 * cm,
                  rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                  title="KAESA Sustainability Briefing %s" % DATE,
                  author="Esther").build(story)
print("PDF saved:", BASE + ".pdf")
