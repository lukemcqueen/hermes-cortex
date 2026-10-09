#!/usr/bin/env python3
"""Generate .docx and .pdf for the daily sustainability briefing.

Usage: ./.venv-brief/bin/python gen_briefing.py 2026-10-10

Reads ~/.hermes/cron/output/sustainability-briefing-<date>.md and writes
sibling .docx and .pdf. Markdown subset handled: H1/H2, bullets, bold
spans, horizontal rules, plain paragraphs — which is all the brief uses.
"""
import os
import re
import sys

OUT = os.path.expanduser("~/.hermes/cron/output")


def parse_md(text: str):
    """Yield (kind, payload) blocks: h1, h2, bullet, para, rule."""
    for raw in text.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.strip() == "---":
            yield ("rule", "")
        elif line.startswith("# "):
            yield ("h1", line[2:].strip())
        elif line.startswith("## "):
            yield ("h2", line[3:].strip())
        elif line.lstrip().startswith("- "):
            yield ("bullet", line.lstrip()[2:].strip())
        else:
            yield ("para", line.strip())


def inline_runs(text: str):
    """Split on **bold** markers -> [(chunk, is_bold)]."""
    parts = re.split(r"\*\*(.+?)\*\*", text)
    runs, bold = [], False
    for p in parts:
        if p:
            runs.append((p.strip(), bold))
        bold = not bold
    return runs or [("", False)]


def linkify(text: str) -> str:
    """Strip markdown link syntax, leaving 'label (url)' readable."""
    return re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)


def build_docx(blocks, path):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for kind, payload in blocks:
        payload = linkify(payload)
        if kind == "h1":
            doc.add_heading(payload, level=0)
        elif kind == "h2":
            doc.add_heading(payload, level=1)
        elif kind == "rule":
            doc.add_paragraph("—" * 30)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            for chunk, is_bold in inline_runs(payload):
                r = p.add_run(chunk)
                r.bold = is_bold
        else:
            p = doc.add_paragraph()
            for chunk, is_bold in inline_runs(payload):
                r = p.add_run(chunk)
                r.bold = is_bold
    doc.save(path)


def build_pdf(blocks, path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (HRFlowable, Paragraph, SimpleDocTemplate,
                                    Spacer)

    base = getSampleStyleSheet()
    h1 = ParagraphStyle("B1", parent=base["Title"], fontSize=17, leading=21,
                        spaceAfter=6, textColor=colors.HexColor("#1a4d3a"))
    h2 = ParagraphStyle("B2", parent=base["Heading2"], fontSize=13, leading=17,
                        spaceBefore=12, spaceAfter=5,
                        textColor=colors.HexColor("#1a4d3a"))
    body = ParagraphStyle("BB", parent=base["BodyText"], fontSize=9.6,
                          leading=13.4, spaceAfter=6)
    bull = ParagraphStyle("BL", parent=body, leftIndent=11, bulletIndent=2,
                          spaceAfter=6)

    doc = SimpleDocTemplate(path, pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title="Sustainable Materials & Market Intelligence Briefing")
    flow = []
    for kind, payload in blocks:
        txt = linkify(payload)
        # escape XML, then re-apply bold
        safe = (txt.replace("&", "&amp;").replace("<", "&lt;")
                   .replace(">", "&gt;"))
        safe = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)
        if kind == "h1":
            flow.append(Paragraph(safe, h1))
        elif kind == "h2":
            flow.append(Paragraph(safe, h2))
        elif kind == "rule":
            flow.append(Spacer(1, 4))
            flow.append(HRFlowable(width="100%", thickness=0.6,
                                   color=colors.HexColor("#b8c9be")))
            flow.append(Spacer(1, 6))
        elif kind == "bullet":
            flow.append(Paragraph(safe, bull, bulletText="\u2022"))
        else:
            flow.append(Paragraph(safe, body))
    doc.build(flow)


def main(date_str: str) -> int:
    base = os.path.join(OUT, f"sustainability-briefing-{date_str}")
    md = base + ".md"
    if not os.path.exists(md):
        print(f"FAIL: missing {md}", file=sys.stderr)
        return 1
    text = open(md, encoding="utf-8").read()
    blocks = list(parse_md(text))
    build_docx(blocks, base + ".docx")
    build_pdf(blocks, base + ".pdf")
    for ext in (".docx", ".pdf"):
        p = base + ext
        print(f"wrote {p} ({os.path.getsize(p)} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "2026-10-10"))
